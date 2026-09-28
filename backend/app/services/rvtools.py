from __future__ import annotations
import csv
import io
from pathlib import Path
from typing import Any
from openpyxl import load_workbook

from .residency import normalize_classification


def _norm(s: Any) -> str:
    return "" if s is None else str(s).strip()


def _lookup(row: dict[str, Any], aliases: list[str], default: Any = "") -> Any:
    normalized = {str(k).strip().lower(): v for k, v in row.items()}
    for alias in aliases:
        if alias.lower() in normalized:
            return normalized[alias.lower()]
    return default


def _number(value: Any, default: float = 0.0) -> float:
    if value in (None, ""):
        return default
    try:
        return float(str(value).replace(",", "").strip())
    except (ValueError, TypeError):
        return default


def _optional_int(value: Any) -> int | None:
    if value in (None, ""):
        return None
    number = _number(value, -1)
    if number < 0:
        raise ValueError(f"Invalid non-negative number '{value}'")
    return int(number)


def normalize_row(row: dict[str, Any]) -> dict[str, Any]:
    memory_mib = _number(_lookup(row, ["Memory", "Memory MiB", "Memory Size MiB", "Memory MB"]))
    memory_gb_direct = _number(_lookup(row, ["Memory GB", "RAM_GB", "RAM GB"]))
    provisioned_mib = _number(_lookup(row, ["Provisioned MiB", "Provisioned MB", "Capacity MiB", "Disk MiB"]))
    storage_gb_direct = _number(_lookup(row, ["Storage GB", "Storage_GB", "Capacity GB"]))
    snapshot_count = int(_number(_lookup(row, ["Snapshots", "Snapshot Count", "Snapshots count"]), 0))
    nic_count = int(_number(_lookup(row, ["NICs", "NIC Count", "Num NICs", "Network adapters"]), 1))

    name = _norm(_lookup(row, ["Name", "VM", "VM Name", "VMName"]))
    if not name:
        raise ValueError("Row does not contain a VM name")

    app_group = _norm(_lookup(row, ["App Group", "Application", "Folder", "vApp", "Resource Pool"], "Ungrouped")) or "Ungrouped"
    source_network = _norm(_lookup(row, ["Network #1", "Network", "VLAN", "Portgroup", "Port Group"], "Unknown")) or "Unknown"

    return {
        "name": name,
        "cpu": int(_number(_lookup(row, ["CPUs", "vCPU", "CPU", "Num CPU"]), 1)),
        "memory_gb": round(memory_gb_direct or (memory_mib / 1024 if memory_mib else 1), 2),
        "storage_gb": round(storage_gb_direct or (provisioned_mib / 1024 if provisioned_mib else 0), 2),
        "os": _norm(_lookup(row, ["OS according to the configuration file", "OS according to config file", "Guest OS", "OS"], "Unknown")) or "Unknown",
        "source_network": source_network,
        "target_network": _norm(_lookup(row, ["Target Network", "AHV Network"], "")),
        "power_state": _norm(_lookup(row, ["Powerstate", "Power State"], "Unknown")) or "Unknown",
        "snapshots": snapshot_count,
        "nic_count": max(nic_count, 1),
        "criticality": (_norm(_lookup(row, ["Criticality", "Business Criticality"], "Medium")) or "Medium").title(),
        "downtime_minutes": int(_number(_lookup(row, ["Downtime Minutes", "Downtime", "Max Downtime"], 60), 60)),
        "app_group": app_group,
        "owner": _norm(_lookup(row, ["Owner", "Application Owner", "Business Owner"], "")),
        "data_classification": normalize_classification(
            _norm(_lookup(row, ["Data Classification", "Classification", "Information Classification"], ""))
        ),
        "residency": _norm(_lookup(row, ["Residency", "Data Residency", "Hosting Country"], "")).upper(),
        "rpo_minutes": _optional_int(_lookup(row, ["RPO Minutes", "RPO"], None)),
        "rto_minutes": _optional_int(_lookup(row, ["RTO Minutes", "RTO"], None)),
    }


def parse_csv(data: bytes) -> list[dict[str, Any]]:
    text = data.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return [normalize_row(row) for row in reader if any(v not in (None, "") for v in row.values())]


def parse_xlsx(data: bytes) -> list[dict[str, Any]]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = wb["vInfo"] if "vInfo" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(v).strip() if v is not None else "" for v in next(rows)]
    result = []
    for values in rows:
        row = dict(zip(headers, values))
        if any(v not in (None, "") for v in values):
            result.append(normalize_row(row))
    return result


def parse_inventory(filename: str, data: bytes) -> list[dict[str, Any]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return parse_csv(data)
    if suffix in {".xlsx", ".xlsm"}:
        return parse_xlsx(data)
    raise ValueError("Supported formats are .csv, .xlsx and .xlsm")
